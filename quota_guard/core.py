"""Deterministic quota rules and local persistence; never sends platform requests.

Money is stored as decimal strings. A snapshot is a replacement of cumulative
costs for the complete configured period, never a delta to add to previous costs.
"""
from __future__ import annotations

from contextlib import closing, contextmanager
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_FLOOR, localcontext
from hashlib import sha256
import json
import os
from pathlib import Path
import re
import sqlite3
import tempfile
from typing import Iterable
from uuid import uuid4
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

TZ = timezone(timedelta(hours=8))
HEADERS = ("API名/人员名", "个人金额（元）")
SCHEMA = 2


class GuardError(ValueError):
    """A readable validation/recovery error which must not trigger key deletion."""


def money(value, *, allow_skip=False) -> Decimal:
    if value is None or isinstance(value, bool):
        raise GuardError("金额必须是数字。")
    try:
        result = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        raise GuardError("金额必须是普通数字，不能填写公式或文字。") from None
    if not result.is_finite() or result > Decimal("1000000000000"):
        raise GuardError("金额必须是有限且合理的数值。")
    if result < 0 and not (allow_skip and result == -1):
        raise GuardError("金额不能为负数，只有 -1 表示暂不管理。")
    if result.as_tuple().exponent < -18:
        raise GuardError("金额精度最多保留 18 位小数。")
    return result


def clean_name(value) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        raise GuardError("API 名称必须是文本。")
    name = value.strip()
    if len(name) > 200 or any(ord(c) < 32 for c in name):
        raise GuardError("API 名称过长或含控制字符。")
    if name.startswith(("=", "+", "@")) or re.search(r"sk-[A-Za-z0-9]{8,}", name):
        raise GuardError("请填写 API 名称，不要填写公式或完整密钥。")
    return name


def _date(value: str) -> str:
    try:
        result = date.fromisoformat(value)
    except (TypeError, ValueError):
        raise GuardError("日期必须使用 YYYY-MM-DD 格式。") from None
    return result.isoformat()


def _identifier(value: str) -> str:
    if not isinstance(value, str) or not value.strip() or len(value) > 200:
        raise GuardError("缺少有效的身份标识。")
    if value.startswith("sk-") or any(ord(c) < 32 for c in value):
        raise GuardError("身份标识不能是密钥或含控制字符。")
    return value.strip()


def _now() -> str:
    return datetime.now(TZ).isoformat(timespec="seconds")


def _safe_detail(value: str) -> str:
    text = str(value)
    text = re.sub(r"sk-[A-Za-z0-9_*.-]+", "[密钥已隐藏]", text)
    text = re.sub(r"(?i)(bearer\s+)\S+", r"\1[已隐藏]", text)
    text = re.sub(r"(?i)((?:cookie|authorization|password|token)\s*[:=]\s*)[^\s,;]+", r"\1[已隐藏]", text)
    return text[:2000]


@dataclass(frozen=True)
class Person:
    name: str
    amount: Decimal


@dataclass(frozen=True)
class Config:
    period_id: str
    people_count: int
    total_budget: Decimal
    start_date: str
    account_id: str = ""

    @property
    def share(self) -> Decimal:
        with localcontext() as context:
            context.prec = 50
            return (self.total_budget / self.people_count).to_integral_value(rounding=ROUND_FLOOR)


class DataStore:
    CONFIG = "管理设置.json"
    WORKBOOK = "人员额度.xlsx"
    DATABASE = "状态.db"
    MARKER = ".初始化中"

    def __init__(self, folder: Path):
        self.folder = Path(folder).resolve()
        self.config_path = self.folder / self.CONFIG
        self.workbook_path = self.folder / self.WORKBOOK
        self.db_path = self.folder / self.DATABASE

    def has_history(self) -> bool:
        if not self.folder.exists():
            return False
        return any((self.folder / n).exists() for n in
                   (self.CONFIG, self.WORKBOOK, self.DATABASE, self.DATABASE + "-wal",
                    self.DATABASE + "-shm", self.MARKER)) or any(
                       self.folder.glob("运行日志-*.jsonl")) or any(self.folder.glob("运行摘要-*.txt"))

    def initialize(self, people_count: int, total_budget, start_date: str) -> Config:
        if isinstance(people_count, bool) or not isinstance(people_count, int) or not 1 <= people_count <= 10000:
            raise GuardError("人数必须是 1 到 10000 的整数。")
        budget = money(total_budget)
        start = _date(start_date)
        if start > datetime.now(TZ).date().isoformat():
            raise GuardError("统计开始日期不能晚于今天。")
        if self.has_history():
            raise GuardError("此目录已有设置或历史，请打开/恢复原数据，不能重新初始化清零。")
        config = Config(str(uuid4()), people_count, budget, start)
        self.folder.mkdir(parents=True, exist_ok=True)
        marker = self.folder / self.MARKER
        try:
            with marker.open("x", encoding="utf-8") as f:
                f.write(config.period_id)
            with closing(sqlite3.connect(self.db_path)) as db, db:
                db.executescript("""
                    CREATE TABLE meta(key TEXT PRIMARY KEY, value TEXT NOT NULL);
                    CREATE TABLE bindings(tracking_id TEXT PRIMARY KEY, person TEXT NOT NULL,
                        active INTEGER NOT NULL DEFAULT 1);
                    CREATE TABLE costs(tracking_id TEXT PRIMARY KEY, amount TEXT NOT NULL,
                        source_id TEXT NOT NULL);
                    CREATE TABLE snapshots(source_id TEXT PRIMARY KEY, digest TEXT NOT NULL,
                        start_date TEXT NOT NULL, end_date TEXT NOT NULL, created TEXT NOT NULL);
                    CREATE TABLE events(id INTEGER PRIMARY KEY AUTOINCREMENT,
                        dedup TEXT UNIQUE, created TEXT NOT NULL, kind TEXT NOT NULL,
                        payload TEXT NOT NULL);
                """)
                for key, value in self._config_dict(config).items():
                    db.execute("INSERT INTO meta VALUES (?,?)", (key, str(value)))
                self._event(db, "initialized", {"people_count": people_count,
                    "total_budget": str(budget), "share": str(config.share), "start_date": start})
            self._write_json(self.config_path, self._config_dict(config))
            self._write_workbook([Person("", Decimal(-1)) for _ in range(people_count)])
            marker.unlink()
            self.export_logs()
        except (OSError, sqlite3.Error) as exc:
            raise GuardError("初始化未完整完成，请保留目录并恢复，不能重建覆盖。" + _safe_detail(exc)) from exc
        return config

    @staticmethod
    def _config_dict(config: Config) -> dict:
        return {"schema": SCHEMA, "period_id": config.period_id,
                "people_count": config.people_count, "total_budget": str(config.total_budget),
                "start_date": config.start_date, "account_id": config.account_id,
                "comparison": ">=", "share_rounding": "floor_yuan", "timezone": "UTC+08:00"}

    def load(self) -> Config:
        if (self.folder / self.MARKER).exists():
            raise GuardError("检测到上次初始化中断，请恢复数据，不会自动清零。")
        if not all(p.is_file() for p in (self.config_path, self.workbook_path, self.db_path)):
            raise GuardError("设置、表格或统计文件缺失，请恢复原文件，不会自动新建周期。")
        try:
            raw = json.loads(self.config_path.read_text(encoding="utf-8"))
            if raw.get("schema") != SCHEMA or raw.get("comparison") != ">=" or raw.get("share_rounding") != "floor_yuan":
                raise GuardError("设置版本或规则不匹配，请恢复正确设置。")
            n = raw["people_count"]
            if isinstance(n, bool) or not isinstance(n, int) or not 1 <= n <= 10000:
                raise GuardError("设置中的人数无效。")
            legacy_account = raw.get("account_id", "")
            if legacy_account != "":
                legacy_account = _identifier(legacy_account)
            config = Config(_identifier(raw["period_id"]), n, money(raw["total_budget"]),
                            _date(raw["start_date"]))
            with self._connect() as db:
                if db.execute("PRAGMA quick_check").fetchone()[0] != "ok":
                    raise GuardError("统计库损坏，请恢复备份。")
                stored = dict(db.execute("SELECT key,value FROM meta"))
                for key, value in self._config_dict(config).items():
                    if key == "account_id":
                        continue
                    if stored.get(key) != str(value):
                        raise GuardError("设置与统计库不匹配，请恢复同一周期的一套文件。")
                # Account identity changes exactly once, in the same SQLite
                # transaction as the first bindings. Old nonempty JSON copies
                # remain assertions; a blank JSON field is not a second store.
                account = stored["account_id"]
                if account != "":
                    account = _identifier(account)
                if legacy_account and legacy_account != account:
                    raise GuardError("设置与统计库的账户不匹配，请恢复同一周期的一套文件。")
                config = Config(config.period_id, n, config.total_budget, config.start_date, account)
            return config
        except (OSError, ValueError, KeyError, TypeError, sqlite3.Error) as exc:
            if isinstance(exc, GuardError):
                raise
            raise GuardError("无法读取已有数据，请恢复文件，不会自动清零。") from exc

    @contextmanager
    def _connect(self):
        # URI read/write prevents SQLite from creating an empty replacement DB.
        try:
            db = sqlite3.connect(self.db_path.as_uri() + "?mode=rw", uri=True, timeout=10)
            db.execute("PRAGMA foreign_keys=ON")
            db.execute("PRAGMA synchronous=FULL")
            with db:
                yield db
        except sqlite3.Error as exc:
            raise GuardError("统计库读取或写入失败，请检查目录权限及文件完整性。") from exc
        finally:
            if "db" in locals():
                db.close()

    @staticmethod
    def _write_json(path: Path, value):
        data = json.dumps(value, ensure_ascii=False, indent=2)
        DataStore._atomic_text(path, data)

    @staticmethod
    def _atomic_text(path: Path, content: str):
        fd, temp = tempfile.mkstemp(prefix=".quota-", suffix=".tmp", dir=path.parent)
        try:
            with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as f:
                f.write(content)
                f.flush()
                os.fsync(f.fileno())
            os.replace(temp, path)
        finally:
            if os.path.exists(temp):
                os.unlink(temp)

    def workbook_hash(self) -> str:
        try:
            return sha256(self.workbook_path.read_bytes()).hexdigest()
        except OSError as exc:
            raise GuardError("无法读取表格，请先保存并关闭占用文件的软件。") from exc

    def _validate_people(self, people: Iterable[Person], count: int) -> list[Person]:
        rows, names = [], set()
        for person in people:
            name, amount = clean_name(person.name), money(person.amount, allow_skip=True)
            if not name and amount != -1:
                raise GuardError("金额为 0 或正数时必须填写 API 名称。")
            if name in names:
                raise GuardError(f"API 名称重复：{name}")
            if name:
                names.add(name)
            rows.append(Person(name, amount))
        if len(names) > count:
            raise GuardError("已填写人员数超过分配人数，请检查表格。")
        return rows

    def read_people(self) -> list[Person]:
        config = self.load()
        if self.workbook_path.stat().st_size > 10 * 1024 * 1024:
            raise GuardError("表格过大，请使用程序生成的两列表格。")
        try:
            with ZipFile(self.workbook_path) as archive:
                entries = archive.infolist()
                if len(entries) > 1000 or sum(e.file_size for e in entries) > 64 * 1024 * 1024:
                    raise GuardError("表格解压内容过大，请使用程序生成的两列表格。")
            wb = load_workbook(self.workbook_path, read_only=True, data_only=False, keep_links=False)
            try:
                if wb.sheetnames != ["人员额度"]:
                    raise GuardError("表格只应有一个名为“人员额度”的工作表。")
                ws = wb["人员额度"]
                if ws.max_row > 20000 or ws.max_column > 2:
                    raise GuardError("表格只能有规定的两列，且行数不能超过 20000。")
                values = ws.iter_rows()
                header = next(values, ())
                if tuple(cell.value for cell in header) != HEADERS:
                    raise GuardError("表头必须为“API名/人员名”和“个人金额（元）”。")
                people = []
                for number, row in enumerate(values, 2):
                    cells = list(row)
                    if any(c.data_type in ("f", "e") for c in cells):
                        raise GuardError(f"第 {number} 行包含公式或错误值，请填写普通数字。")
                    name, amount = cells[0].value, cells[1].value
                    if name is None and amount is None:
                        continue
                    try:
                        people.append(Person(clean_name(name), money(amount, allow_skip=True)))
                    except GuardError as exc:
                        raise GuardError(f"第 {number} 行：{exc}") from exc
                return self._validate_people(people, config.people_count)
            finally:
                wb.close()
        except GuardError:
            raise
        except Exception as exc:
            raise GuardError("表格无法读取，请确认是完整保存的 .xlsx 文件。") from exc

    def _write_workbook(self, people: list[Person]):
        wb = Workbook()
        ws = wb.active
        ws.title = "人员额度"
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="245B78")
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 24
        ws.freeze_panes = "A2"
        for person in people:
            # Numeric text preserves decimal precision beyond Excel's 15 digits.
            ws.append([person.name, str(person.amount)])
        validation = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="-1")
        validation.promptTitle = "金额说明"
        validation.prompt = "-1 暂不管理；0 停用；正数为个人累计额度。"
        validation.showInputMessage = True
        ws.add_data_validation(validation)
        validation.add(f"B2:B{max(len(people)+1, 2)}")
        fd, temp = tempfile.mkstemp(prefix=".quota-", suffix=".xlsx", dir=self.folder)
        os.close(fd)
        try:
            wb.save(temp)
            os.replace(temp, self.workbook_path)
        finally:
            wb.close()
            if os.path.exists(temp):
                os.unlink(temp)

    def save_people(self, people: list[Person], expected_hash: str | None = None):
        config = self.load()
        rows = self._validate_people(people, config.people_count)
        if expected_hash is not None and expected_hash != self.workbook_hash():
            raise GuardError("表格已被其他软件修改，请重新读取后再保存。")
        try:
            self._write_workbook(rows)
        except OSError as exc:
            raise GuardError("表格无法保存，请关闭占用它的 Excel/WPS 后重试。") from exc
        with self._connect() as db:
            self._event(db, "people_saved", {"workbook_hash": self.workbook_hash(),
                "people": [{"name": p.name, "amount": str(p.amount)} for p in rows if p.name]})
        self.evaluate()

    def active_bindings(self) -> list[dict]:
        """Persisted identities awaiting verified absence, including after restart.

        This is an expected set for reconciliation, not proof that a Key still
        exists. Callers must verify a surviving identity before retiring one.
        """
        self.load()
        with self._connect() as db:
            return [{"name": name, "tracking_id": key} for key, name in db.execute(
                "SELECT tracking_id,person FROM bindings WHERE active=1 ORDER BY tracking_id")]

    def active_tracking_ids(self) -> set[str]:
        """Return identities not yet retired, without assuming they still exist."""
        return {row["tracking_id"] for row in self.active_bindings()}

    def latest_action(self, tracking_id: str) -> str | None:
        """Latest persisted action status, across the complete event history."""
        self.load()
        key_id = _identifier(tracking_id)
        with self._connect() as db:
            for (payload,) in db.execute("SELECT payload FROM events WHERE kind='action' ORDER BY id DESC"):
                event = json.loads(payload)
                if event.get("tracking_id") == key_id:
                    return event["status"]
        return None

    def bind_keys(self, keys: list[dict], account_id: str):
        config = self.load()
        account = _identifier(account_id)
        if config.account_id and account != config.account_id:
            raise GuardError("登录账户与此数据目录不一致，拒绝绑定。")
        people = [p for p in self.read_people() if p.amount >= 0]
        by_name: dict[str, list[str]] = {}
        seen_ids = set()
        for key in keys:
            name, key_id = clean_name(key["name"]), _identifier(key["tracking_id"])
            if key_id in seen_ids:
                raise GuardError("平台返回重复 Tracking ID，无法安全绑定。")
            seen_ids.add(key_id)
            by_name.setdefault(name, []).append(key_id)
        with self._connect() as db:
            db.execute("BEGIN IMMEDIATE")
            bindings = {r[0]: r[1] for r in db.execute("SELECT tracking_id, person FROM bindings")}
            active = {r[0]: r[1] for r in db.execute("SELECT person,tracking_id FROM bindings WHERE active=1")}
            additions = []
            for p in people:
                matches = by_name.get(p.name, [])
                historical = [key for key, name in bindings.items() if name == p.name]
                if historical and p.name not in active:
                    if matches:
                        raise GuardError(f"{p.name} 已有退休 Key 的历史，出现同名 Key 必须人工核对换绑，不能自动清零。")
                    continue
                if len(matches) != 1:
                    raise GuardError(f"{p.name}：平台必须存在且仅存在一个同名 Key。")
                key_id = matches[0]
                if p.name in active and active[p.name] != key_id:
                    raise GuardError(f"{p.name} 的 Tracking ID 已变更，必须人工核对历史并换绑，不能自动清零。")
                if key_id in bindings and bindings[key_id] != p.name:
                    raise GuardError(f"{p.name} 的 Key 已绑定其他人员，拒绝按新名字覆盖历史。")
                if key_id not in bindings:
                    additions.append((key_id, p.name))
            for key_id, name in additions:
                db.execute("INSERT INTO bindings(tracking_id,person) VALUES (?,?)", (key_id, name))
            if not config.account_id:
                db.execute("UPDATE meta SET value=? WHERE key='account_id'", (account,))
            if additions or not config.account_id:
                self._event(db, "keys_bound", {"bindings": [{"name": n, "tracking_id": k} for k, n in additions]})
        self.evaluate()

    def apply_snapshot(self, totals: dict[str, Decimal], source_id: str,
                       start_date: str, end_date: str) -> list[dict]:
        config = self.load()
        start, end = _date(start_date), _date(end_date)
        if start != config.start_date or end < start or end > datetime.now(TZ).date().isoformat():
            raise GuardError("账单时间范围错误，必须从本期起点开始且不晚于今天。")
        if not config.account_id:
            raise GuardError("请先绑定已核实的平台账户及人员 Key。")
        source = _identifier(source_id)
        normalized = {_identifier(k): money(v) for k, v in totals.items()}
        if len(normalized) != len(totals):
            raise GuardError("账单存在重复的 Tracking ID，请核对来源。")
        digest = sha256(json.dumps({"start": start, "end": end,
            "totals": {k: format(v, "f") for k, v in sorted(normalized.items())}}, sort_keys=True).encode()).hexdigest()
        people = self.read_people()
        with self._connect() as db:
            db.execute("BEGIN IMMEDIATE")
            existing = db.execute("SELECT digest FROM snapshots WHERE source_id=?", (source,)).fetchone()
            if existing:
                if existing[0] != digest:
                    raise GuardError("相同账单标识对应不同内容，请重新核实来源。")
            else:
                previous_end = db.execute("SELECT MAX(end_date) FROM snapshots").fetchone()[0]
                if previous_end and end < previous_end:
                    raise GuardError("导入的账单截止日期早于现有统计，拒绝回退。")
                required = {r[0] for r in db.execute("SELECT tracking_id FROM bindings")}
                received = required.intersection(normalized)
                if not received:
                    raise GuardError("账单没有任何已绑定 Key 的可核验费用，未覆盖原统计。")
                old = {k: Decimal(v) for k, v in db.execute("SELECT tracking_id,amount FROM costs")}
                for key_id in received:
                    if key_id in old and normalized[key_id] < old[key_id]:
                        raise GuardError("累计费用下降，请核对账单修订或日期范围，未覆盖原统计。")
                for key_id in received:
                    db.execute("INSERT INTO costs VALUES (?,?,?) ON CONFLICT(tracking_id) DO UPDATE "
                               "SET amount=excluded.amount, source_id=excluded.source_id",
                               (key_id, str(normalized[key_id]), source))
                db.execute("INSERT INTO snapshots VALUES (?,?,?,?,?)", (source, digest, start, end, _now()))
                self._event(db, "snapshot", {"source_id": source, "start_date": start,
                    "end_date": end, "tracked_key_count": len(received),
                    "missing_key_count": len(required - received)})
            # Costs and their trigger events commit together. Replaying an old
            # source still reconciles events against the current workbook.
            results = self._evaluate(db, config, people)
            self._record_triggers(db, config, results)
        self.export_logs()
        return results

    def evaluate(self) -> list[dict]:
        """Evaluate current rules and idempotently reconcile their audit events.

        Called on startup and after workbook edits as well as after collection,
        so zero budgets and external edits do not depend on a new cost report.
        """
        config, people = self.load(), self.read_people()
        with self._connect() as db:
            db.execute("BEGIN IMMEDIATE")
            results = self._evaluate(db, config, people)
            self._record_triggers(db, config, results)
        self.export_logs()
        return results

    def _record_triggers(self, db, config: Config, results: list[dict]):
        for result in results:
            if result["status"] != "triggered":
                continue
            payload = {k: str(v) if isinstance(v, Decimal) else v for k, v in result.items()}
            dedup = sha256(json.dumps({"period": config.period_id, "name": result["name"],
                "key": result["tracking_id"], "amount": str(result["amount"]),
                "reasons": result["reasons"]}, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
            self._event(db, "quota_triggered", payload, dedup="trigger:"+dedup)

    @staticmethod
    def _evaluate(db, config: Config, people: list[Person]) -> list[dict]:
        bindings = list(db.execute("SELECT tracking_id,person,active FROM bindings"))
        costs = {k: (Decimal(v), source) for k, v, source in
                 db.execute("SELECT tracking_id,amount,source_id FROM costs")}
        latest = db.execute("SELECT source_id FROM snapshots ORDER BY rowid DESC LIMIT 1").fetchone()
        latest_source = latest[0] if latest else None
        results = []
        for person in people:
            ids = [key for key, name, _ in bindings if name == person.name]
            active_ids = [key for key, name, active in bindings if name == person.name and active]
            with localcontext() as context:
                context.prec = 50
                known = [costs[k][0] for k in ids if k in costs]
                spent = sum(known, Decimal(0)) if known else None
            complete = bool(ids) and latest_source is not None and all(
                k in costs and costs[k][1] == latest_source for k in ids)
            result = {"name": person.name, "amount": person.amount, "spent": spent,
                      "share": config.share, "tracking_id": active_ids[0] if len(active_ids) == 1 else "",
                      "status": "ok", "reasons": [], "complete": complete,
                      "spent_is_lower_bound": spent is not None and not complete}
            if person.amount == -1:
                result["status"] = "skipped"
            elif ids and not active_ids:
                result["status"] = "retired"
                result["reasons"].append("已核实平台列表中不存在，保留历史费用；不据此认定删除原因")
            elif person.amount == 0 or config.share == 0:
                result["status"] = "triggered"
                if person.amount == 0:
                    result["reasons"].append("个人额度为零")
                if config.share == 0:
                    result["reasons"].append("公共人均额度为零")
            elif spent is None or not result["tracking_id"]:
                result["status"] = "unknown"
                result["reasons"].append("尚未完成 Key 绑定或缺少完整周期账单")
            else:
                if spent >= person.amount:
                    result["reasons"].append("达到个人额度")
                if spent >= config.share:
                    result["reasons"].append("达到公共人均额度")
                if result["reasons"]:
                    result["status"] = "triggered"
                elif not complete:
                    result["status"] = "unknown"
                    result["reasons"].append("最近费用快照缺少该人员部分记录，历史金额仅为已知下界")
            results.append(result)
        return results

    def record_action(self, name: str, tracking_id: str, status: str, detail: str):
        self.load()
        name, key_id = clean_name(name), _identifier(tracking_id)
        if status not in {"pending", "prepared", "unknown_result", "absent_verified",
                          "requested", "confirmed", "failed", "unknown", "cancelled"}:
            raise GuardError("无效的操作状态。")
        with self._connect() as db:
            db.execute("BEGIN IMMEDIATE")
            bound = db.execute("SELECT person FROM bindings WHERE tracking_id=?", (key_id,)).fetchone()
            if not bound or bound[0] != name:
                raise GuardError("操作目标与既有绑定不一致，拒绝记录成功。")
            if status == "absent_verified":
                db.execute("UPDATE bindings SET active=0 WHERE tracking_id=?", (key_id,))
            self._event(db, "action", {"name": name, "tracking_id": key_id,
                "status": status, "detail": _safe_detail(detail)})
        self.export_logs()

    @staticmethod
    def _event(db, kind: str, payload: dict, dedup: str | None = None):
        db.execute("INSERT OR IGNORE INTO events(dedup,created,kind,payload) VALUES (?,?,?,?)",
                   (dedup, _now(), kind, json.dumps(payload, ensure_ascii=False)))

    def events(self, limit: int = 200) -> list[dict]:
        self.load()
        with self._connect() as db:
            rows = db.execute("SELECT id,created,kind,payload FROM events ORDER BY id DESC LIMIT ?",
                              (max(1, min(int(limit), 10000)),)).fetchall()
        return [{"id": i, "created": t, "kind": k, **json.loads(p)} for i, t, k, p in rows]

    def export_logs(self):
        """SQLite is authoritative. Rebuilding daily files makes crash recovery idempotent."""
        with self._connect() as db:
            rows = db.execute("SELECT id,created,kind,payload FROM events ORDER BY id").fetchall()
        daily: dict[str, list[dict]] = {}
        for number, created, kind, payload in rows:
            daily.setdefault(created[:10], []).append({"id": number, "created": created,
                                                       "kind": kind, **json.loads(payload)})
        try:
            for day, events in daily.items():
                jsonl = "".join(json.dumps(e, ensure_ascii=False) + "\n" for e in events)
                readable = "".join(f"#{e['id']} {e['created']} | {e['kind']} | " +
                    json.dumps({k: v for k, v in e.items() if k not in {"id", "created", "kind"}}, ensure_ascii=False)
                    + "\n" for e in events)
                self._atomic_text(self.folder / f"运行日志-{day}.jsonl", jsonl)
                self._atomic_text(self.folder / f"运行摘要-{day}.txt", readable)
        except OSError as exc:
            raise GuardError("事件已保存在统计库，但同目录日志写入失败；请停止实际操作并恢复目录权限。") from exc
