from contextlib import closing
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import json
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from quota_guard.core import Config, DataStore, GuardError, Person, TZ, money


class CoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.store = DataStore(self.folder)
        self.start = datetime.now(TZ).date().isoformat()
        self.store.initialize(14, "1500", self.start)

    def tearDown(self):
        self.temp.cleanup()

    def people(self, *pairs):
        self.store.save_people([Person(name, Decimal(str(amount))) for name, amount in pairs])

    def bind(self, *names):
        self.store.bind_keys([{"name": n, "tracking_id": f"id-{n}"} for n in names], "account-test")

    def snapshot(self, values, sid="sample"):
        return self.store.apply_snapshot({f"id-{name}": Decimal(str(amount)) for name, amount in values.items()},
                                         sid, self.start, self.start)

    def test_init_template_and_fixed_share(self):
        self.assertEqual(self.store.load().share, Decimal(107))
        self.assertEqual(len(self.store.read_people()), 14)
        self.assertTrue(all(p == Person("", Decimal(-1)) for p in self.store.read_people()))
        self.people(("甲", 80), ("乙", 200))
        self.assertEqual(self.store.load().share, Decimal(107))

    def test_boundaries_and_skipping(self):
        self.people(("甲", 80), ("乙", 200), ("丙", -1), ("丁", 0))
        self.bind("甲", "乙", "丁")
        before = {r["name"]: r for r in self.store.evaluate()}
        self.assertEqual(before["甲"]["status"], "unknown")
        self.assertEqual(before["丙"]["status"], "skipped")
        self.assertEqual(before["丁"]["status"], "triggered")
        rows = {r["name"]: r for r in self.snapshot({"甲":80, "乙":107, "丁":0})}
        self.assertEqual(rows["甲"]["reasons"], ["达到个人额度"])
        self.assertEqual(rows["乙"]["reasons"], ["达到公共人均额度"])
        self.assertEqual(rows["丁"]["status"], "triggered")

    def test_jump_over_and_precision(self):
        self.people(("甲", 80))
        self.bind("甲")
        self.assertEqual(self.snapshot({"甲":"79.999999"})[0]["status"], "ok")
        self.assertEqual(self.snapshot({"甲":"80.000001"}, "new")[0]["status"], "triggered")

    def test_snapshot_is_replacement_and_idempotent(self):
        self.people(("甲", 80))
        self.bind("甲")
        self.snapshot({"甲":90})
        events_before = len(self.store.events())
        self.snapshot({"甲":90})
        self.assertEqual(len(self.store.events()), events_before)
        self.snapshot({"甲":95}, "sample-new")
        self.assertEqual(self.store.evaluate()[0]["spent"], Decimal(95))
        triggers = [e for e in self.store.events() if e["kind"] == "quota_triggered"]
        self.assertEqual(len(triggers), 1)
        self.assertEqual(DataStore(self.folder).evaluate()[0]["spent"], Decimal(95))

    def test_missing_decreasing_and_conflicting_data_rejected(self):
        self.people(("甲", 80))
        self.bind("甲")
        self.snapshot({"甲":50})
        for totals, source in [({}, "missing"), ({"甲":49}, "lower"), ({"甲":51}, "sample")]:
            with self.assertRaises(GuardError):
                self.snapshot(totals, source)
        self.assertEqual(self.store.evaluate()[0]["spent"], Decimal(50))

    def test_wrong_period_rejected(self):
        self.people(("甲", 80))
        self.bind("甲")
        with self.assertRaises(GuardError):
            self.store.apply_snapshot({"id-甲":Decimal(20)}, "bad", "2000-01-01", self.start)

    def test_key_rename_or_replacement_cannot_reset_history(self):
        self.people(("甲", 80))
        self.bind("甲")
        self.snapshot({"甲":70})
        with self.assertRaises(GuardError):
            self.store.bind_keys([{"name":"甲", "tracking_id":"new-id"}], "account-test")
        self.people(("乙", 80))
        with self.assertRaises(GuardError):
            self.store.bind_keys([{"name":"乙", "tracking_id":"id-甲"}], "account-test")
        self.people(("甲", 80))
        self.assertEqual(self.store.evaluate()[0]["spent"], Decimal(70))

    def test_duplicate_platform_names_and_different_account_rejected(self):
        self.people(("甲", 80))
        with self.assertRaises(GuardError):
            self.store.bind_keys([{"name":"甲", "tracking_id":"one"}, {"name":"甲", "tracking_id":"two"}], "account-test")
        self.bind("甲")
        with self.assertRaises(GuardError):
            self.store.bind_keys([{"name":"甲", "tracking_id":"id-甲"}], "wrong-account")

    def test_missing_file_does_not_reinitialize(self):
        self.store.workbook_path.unlink()
        with self.assertRaises(GuardError): self.store.load()
        with self.assertRaises(GuardError): self.store.initialize(14, "1500", self.start)
        self.assertTrue(self.store.db_path.exists())

    def test_settings_tampering_is_rejected(self):
        content = self.store.config_path.read_text(encoding="utf-8")
        self.store.config_path.write_text(content.replace('"1500"', '"2000"'), encoding="utf-8")
        with self.assertRaises(GuardError): self.store.load()

    def test_partial_initialization_is_not_fresh(self):
        (self.folder / self.store.MARKER).write_text("interrupted", encoding="utf-8")
        with self.assertRaises(GuardError): self.store.load()
        with self.assertRaises(GuardError): self.store.initialize(14, "1500", self.start)

    def test_invalid_values_and_duplicates(self):
        for value in ("NaN", "Infinity", "-2", "-0.5", "=1+1", True, None):
            with self.subTest(value=value), self.assertRaises(GuardError):
                money(value, allow_skip=True)
        with self.assertRaises(GuardError): self.people(("甲", 1), ("甲", 2))
        with self.assertRaises(GuardError): self.people(("", 0))

    def test_excel_formulas_rejected(self):
        wb = load_workbook(self.store.workbook_path)
        wb.active["A2"] = "甲"
        wb.active["B2"] = "=80"
        wb.save(self.store.workbook_path)
        wb.close()
        with self.assertRaises(GuardError): self.store.read_people()

    def test_concurrent_edit_protection(self):
        fingerprint = self.store.workbook_hash()
        self.people(("甲", 80))
        with self.assertRaises(GuardError):
            self.store.save_people([Person("乙", Decimal(10))], expected_hash=fingerprint)
        self.assertEqual(self.store.read_people()[0].name, "甲")

    def test_action_results_logged_and_secrets_redacted(self):
        self.people(("甲", 0))
        self.bind("甲")
        self.store.record_action("甲", "id-甲", "pending", "准备执行")
        self.store.record_action("甲", "id-甲", "unknown", "Bearer sensitive123 sk-testlongsecret cookie=sessionsecret")
        logs = "\n".join(p.read_text(encoding="utf-8") for p in self.folder.glob("运行*.*"))
        self.assertIn("unknown", logs)
        for secret in ("sensitive123", "sk-testlongsecret", "sessionsecret"):
            self.assertNotIn(secret, logs)
        with self.assertRaises(GuardError):
            self.store.record_action("乙", "id-甲", "confirmed", "wrong target")

    def test_zero_shared_budget_triggers_active_people(self):
        other = DataStore(self.folder / "other")
        other.initialize(14, 0, self.start)
        other.save_people([Person("甲", Decimal(50)), Person("乙", Decimal(-1))])
        self.assertEqual([r["status"] for r in other.evaluate()], ["triggered", "skipped"])

    def test_partial_snapshot_updates_present_keys_without_inventing_zero(self):
        self.people(("甲", 80), ("乙", 80))
        self.bind("甲", "乙")
        rows = {r["name"]: r for r in self.snapshot({"甲": 81})}
        self.assertEqual(rows["甲"]["status"], "triggered")
        self.assertTrue(rows["甲"]["complete"])
        self.assertIsNone(rows["乙"]["spent"])
        self.assertEqual(rows["乙"]["status"], "unknown")
        self.assertFalse(rows["乙"]["complete"])

    def test_missing_latest_cost_keeps_history_as_lower_bound(self):
        self.people(("甲", 80), ("乙", 80))
        self.bind("甲", "乙")
        self.snapshot({"甲": 20, "乙": 30})
        rows = {r["name"]: r for r in self.snapshot({"甲": 25}, "partial")}
        self.assertEqual(rows["乙"]["spent"], Decimal(30))
        self.assertEqual(rows["乙"]["status"], "unknown")
        self.assertTrue(rows["乙"]["spent_is_lower_bound"])
        self.assertEqual(rows["甲"]["status"], "ok")
        self.assertFalse(rows["甲"]["spent_is_lower_bound"])
        # Replaying older data must not make an omitted key look current.
        self.snapshot({"甲": 20, "乙": 30})
        replayed = {r["name"]: r for r in self.store.evaluate()}
        self.assertEqual(replayed["乙"]["status"], "unknown")
        rows = {r["name"]: r for r in self.snapshot({"甲": 25, "乙": 35}, "complete")}
        self.assertEqual(rows["乙"]["status"], "ok")
        self.assertTrue(rows["乙"]["complete"])

    def test_known_lower_bound_still_triggers(self):
        self.people(("甲", 80), ("乙", 80))
        self.bind("甲", "乙")
        self.snapshot({"甲": 20, "乙": 80})
        rows = {r["name"]: r for r in self.snapshot({"甲": 25}, "partial")}
        self.assertEqual(rows["乙"]["status"], "triggered")
        self.assertEqual(rows["乙"]["spent"], Decimal(80))
        self.assertTrue(rows["乙"]["spent_is_lower_bound"])

    def test_verified_absence_retires_without_losing_costs(self):
        self.people(("甲", 80), ("乙", 80))
        self.bind("甲", "乙")
        self.snapshot({"甲": 80, "乙": 20})
        self.store.record_action("甲", "id-甲", "absent_verified", "平台列表已不存在")
        self.assertEqual(self.store.active_tracking_ids(), {"id-乙"})
        rows = {r["name"]: r for r in self.store.evaluate()}
        self.assertEqual(rows["甲"]["status"], "retired")
        self.assertEqual(rows["甲"]["tracking_id"], "")
        self.assertEqual(rows["甲"]["spent"], Decimal(80))
        self.bind("乙")
        rows = {r["name"]: r for r in self.snapshot({"乙": 30}, "after-retirement")}
        self.assertEqual(rows["甲"]["status"], "retired")
        self.assertEqual(rows["甲"]["spent"], Decimal(80))
        self.assertEqual(rows["乙"]["status"], "ok")
        with self.assertRaises(GuardError):
            self.store.bind_keys([{"name": "甲", "tracking_id": "replacement"},
                                  {"name": "乙", "tracking_id": "id-乙"}], "account-test")
        with self.assertRaises(GuardError):
            self.bind("甲", "乙")

    def test_action_contract_and_uncertain_results_do_not_retire(self):
        self.people(("甲", 0))
        self.bind("甲")
        for state in ("pending", "prepared", "unknown_result", "requested", "confirmed", "failed", "unknown", "cancelled"):
            self.store.record_action("甲", "id-甲", state, state)
            self.assertEqual(self.store.active_tracking_ids(), {"id-甲"})
        with self.assertRaises(GuardError):
            self.store.record_action("乙", "id-甲", "absent_verified", "wrong target")
        self.assertEqual(self.store.active_tracking_ids(), {"id-甲"})
        self.store.record_action("甲", "id-甲", "absent_verified", "平台列表已不存在")
        self.assertEqual(self.store.evaluate()[0]["status"], "retired")
        self.assertEqual(self.store.active_tracking_ids(), set())
        self.people(("甲", -1))
        self.assertEqual(self.store.evaluate()[0]["status"], "skipped")

    def test_first_binding_uses_one_persistent_account_source(self):
        self.people(("甲", 80))
        original_settings = self.store.config_path.read_bytes()
        # Binding must not depend on a second filesystem commit after SQLite.
        with patch.object(self.store, "_write_json", side_effect=OSError("simulated write failure")) as writer:
            self.bind("甲")
        writer.assert_not_called()
        self.assertEqual(self.store.config_path.read_bytes(), original_settings)
        restored = DataStore(self.folder)
        self.assertEqual(restored.load().account_id, "account-test")
        self.assertEqual(restored.active_tracking_ids(), {"id-甲"})
        with self.assertRaises(GuardError):
            restored.bind_keys([{"name": "甲", "tracking_id": "id-甲"}], "other-account")

    def test_legacy_account_copy_must_match_authoritative_database(self):
        self.people(("甲", 80))
        self.bind("甲")
        config = json.loads(self.store.config_path.read_text(encoding="utf-8"))
        config["account_id"] = "account-test"
        self.store.config_path.write_text(json.dumps(config), encoding="utf-8")
        self.assertEqual(self.store.load().account_id, "account-test")
        config["account_id"] = "different-account"
        self.store.config_path.write_text(json.dumps(config), encoding="utf-8")
        with self.assertRaises(GuardError):
            self.store.load()
        config["account_id"] = "account-test"
        self.store.config_path.write_text(json.dumps(config), encoding="utf-8")
        with closing(sqlite3.connect(self.store.db_path)) as db, db:
            db.execute("UPDATE meta SET value='' WHERE key='account_id'")
        with self.assertRaises(GuardError):
            self.store.load()

    def test_zero_budget_records_trigger_without_any_snapshot(self):
        self.people(("甲", 0), ("乙", -1))
        self.bind("甲")
        events = [e for e in self.store.events() if e["kind"] == "quota_triggered"]
        bound_events = [e for e in events if e["tracking_id"] == "id-甲"]
        self.assertEqual(len(bound_events), 1)
        self.assertEqual(bound_events[0]["reasons"], ["个人额度为零"])
        self.assertFalse(any(e["kind"] == "snapshot" for e in self.store.events()))
        self.assertFalse(any(e["name"] == "乙" for e in events))
        self.store.evaluate()
        DataStore(self.folder).evaluate()
        self.assertEqual(len([e for e in self.store.events() if e["kind"] == "quota_triggered"]), len(events))

    def test_budget_change_logs_trigger_with_unchanged_cost_report(self):
        self.people(("甲", 100))
        self.bind("甲")
        self.snapshot({"甲": 60})
        self.people(("甲", 50))
        events = [e for e in self.store.events() if e["kind"] == "quota_triggered"]
        self.assertEqual(len(events), 1)
        self.assertEqual(events[0]["amount"], "50")
        self.assertEqual(events[0]["spent"], "60")
        self.snapshot({"甲": 60})
        self.assertEqual(len([e for e in self.store.events() if e["kind"] == "quota_triggered"]), 1)
        # An external Excel edit does not call save_people; a replay must still
        # evaluate and log the new rule without rewriting the existing costs.
        wb = load_workbook(self.store.workbook_path)
        wb.active["B2"] = "40"
        wb.save(self.store.workbook_path)
        wb.close()
        self.snapshot({"甲": 60})
        events = [e for e in self.store.events() if e["kind"] == "quota_triggered"]
        self.assertEqual({e["amount"] for e in events}, {"40", "50"})

    def test_snapshot_and_trigger_commit_or_rollback_together(self):
        self.people(("甲", 80))
        self.bind("甲")
        with patch.object(self.store, "_record_triggers", side_effect=RuntimeError("simulated interruption")):
            with self.assertRaises(RuntimeError):
                self.snapshot({"甲": 90})
        restored = DataStore(self.folder)
        self.assertIsNone(restored.evaluate()[0]["spent"])
        self.assertFalse(any(e["kind"] == "snapshot" for e in restored.events()))
        self.snapshot({"甲": 90})
        self.assertEqual(len([e for e in self.store.events() if e["kind"] == "quota_triggered"]), 1)

    def test_replay_repairs_legacy_missing_trigger_and_export_failure(self):
        self.people(("甲", 80))
        self.bind("甲")
        self.snapshot({"甲": 90})
        # Model an old-version interruption after the snapshot commit but before
        # its separate trigger commit. The unchanged source must repair it.
        with closing(sqlite3.connect(self.store.db_path)) as db, db:
            db.execute("DELETE FROM events WHERE kind='quota_triggered'")
        with patch.object(self.store, "export_logs", side_effect=GuardError("simulated export failure")):
            with self.assertRaises(GuardError):
                self.snapshot({"甲": 90})
        self.snapshot({"甲": 90})
        self.assertEqual(len([e for e in self.store.events() if e["kind"] == "quota_triggered"]), 1)
        exported = [json.loads(line) for path in self.folder.glob("运行日志-*.jsonl")
                    for line in path.read_text(encoding="utf-8").splitlines()]
        self.assertEqual(len([e for e in exported if e["kind"] == "quota_triggered"]), 1)

    def test_pending_bindings_and_latest_action_survive_restart(self):
        self.people(("甲", 80), ("乙", 80))
        self.bind("甲", "乙")
        self.assertIsNone(self.store.latest_action("id-甲"))
        self.store.record_action("甲", "id-甲", "requested", "等待结果")
        self.store.record_action("乙", "id-乙", "pending", "另一个目标")
        self.people(("甲", -1), ("乙", 80))
        restored = DataStore(self.folder)
        self.assertEqual(restored.latest_action("id-甲"), "requested")
        self.assertEqual(restored.latest_action("id-乙"), "pending")
        self.assertEqual({r["tracking_id"] for r in restored.active_bindings()}, {"id-甲", "id-乙"})
        restored.record_action("甲", "id-甲", "unknown_result", "尚未核实")
        self.assertEqual(restored.latest_action("id-甲"), "unknown_result")
        restored.record_action("甲", "id-甲", "absent_verified", "有存活身份核实")
        self.assertEqual(restored.latest_action("id-甲"), "absent_verified")
        self.assertEqual(restored.active_bindings(), [{"name": "乙", "tracking_id": "id-乙"}])
        self.assertIsNone(restored.latest_action("never-bound"))

    def test_orphan_sqlite_sidecars_prevent_reinitialization(self):
        for suffix in ("-wal", "-shm"):
            with self.subTest(suffix=suffix):
                folder = self.folder / suffix
                folder.mkdir()
                artifact = folder / (DataStore.DATABASE + suffix)
                artifact.write_bytes(b"preserved history")
                other = DataStore(folder)
                self.assertTrue(other.has_history())
                with self.assertRaises(GuardError):
                    other.initialize(14, "1500", self.start)
                self.assertEqual(artifact.read_bytes(), b"preserved history")


if __name__ == "__main__":
    unittest.main()
